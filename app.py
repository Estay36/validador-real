"""
Sistema de automatización de revisión de planillas de pago
Módulo: Recepción real de pagos provisorios

Genera la hoja "Recepción PROVISORIOS" y el archivo "GRAL Recepción" a partir de:
  1. Archivo CARGAS MEDIOS (pestañas PROVISORIOS, Monitor Prov,
     Recepción PROVISORIOS y Aprobadores)
  2. Informe Real de Inversión del mes

Ejecutar localmente:   streamlit run app.py
"""
from __future__ import annotations

import datetime as dt
import hashlib
import io

import pandas as pd
import streamlit as st

import config as cfg
from exportador import exportar_gral, exportar_recepcion
from generador import aplicar_ajustes, procesar
from loaders import (
    ErrorEstructura, hojas_informe_real, leer_aprobadores, leer_informe_real,
    leer_monitor_prov, leer_provisorios, leer_recepcion_existente,
)

st.set_page_config(page_title="Recepción real de provisorios", layout="wide")

NOMBRE_MES = {k: v[1].capitalize() for k, v in cfg.MESES.items()}


# ---------------------------------------------------------------------------
# Lectura con caché (el archivo de cargas pesa varios MB)
# ---------------------------------------------------------------------------
def _hash(b: bytes) -> str:
    return hashlib.md5(b).hexdigest()


@st.cache_data(show_spinner=False)
def cargar_archivo_cargas(_contenido: bytes, clave: str):
    """`clave` (hash del archivo) es lo único que usa la caché; el contenido se excluye."""
    prov = leer_provisorios(io.BytesIO(_contenido))
    monitor = leer_monitor_prov(io.BytesIO(_contenido))
    hist = leer_recepcion_existente(io.BytesIO(_contenido))
    aprob = leer_aprobadores(io.BytesIO(_contenido))
    return prov, monitor, hist, aprob


@st.cache_data(show_spinner=False)
def listar_hojas_informe(_contenido: bytes, clave: str):
    return hojas_informe_real(io.BytesIO(_contenido))


@st.cache_data(show_spinner=False)
def cargar_informe(_contenido: bytes, clave: str, hoja: str, mes: int | None):
    inf = leer_informe_real(io.BytesIO(_contenido), hoja=hoja, mes=mes)
    return inf.datos, inf.mes, inf.anio


def formato_pesos(x: float) -> str:
    signo = "-" if x < 0 else ""
    return f"{signo}${abs(x):,.0f}".replace(",", ".")


# ---------------------------------------------------------------------------
# Identificación automática de archivos
# ---------------------------------------------------------------------------
@st.cache_data(show_spinner=False)
def identificar_archivo(_contenido: bytes, clave: str) -> str:
    """
    Reconoce el archivo por sus pestañas, no por su nombre:
      - CARGAS MEDIOS: tiene 'PROVISORIOS' y 'Monitor Prov'
      - INFORME REAL : tiene una hoja 'REAL <MES> <AÑO>' o 'BASES <AÑO>'
      - GRAL         : tiene 'Data' con 'Documento compras' (es una salida)
    """
    from openpyxl import load_workbook
    from utils import buscar_hoja
    try:
        wb = load_workbook(io.BytesIO(_contenido), read_only=True)
    except Exception:
        return "ilegible"
    hojas = wb.sheetnames
    if buscar_hoja(hojas, cfg.HOJA_PROVISORIOS) and buscar_hoja(hojas, cfg.HOJA_MONITOR_PROV):
        return "cargas"
    if hojas_informe_real(io.BytesIO(_contenido)):
        return "informe"
    if buscar_hoja(hojas, "Data"):
        fila = next(wb[buscar_hoja(hojas, "Data")].iter_rows(max_row=1, values_only=True), ())
        if any(str(v).strip() == "Documento compras" for v in fila if v):
            return "gral"
    if buscar_hoja(hojas, "DETALLE ESTIMADO") or buscar_hoja(hojas, "DATA") or buscar_hoja(hojas, "Detalle"):
        return "provisorio"
    return "desconocido"


DESCRIPCION_TIPO = {
    "cargas": ("Archivo de cargas de medios", "Se usa: PROVISORIOS, Monitor Prov, Recepción PROVISORIOS y Aprobadores"),
    "informe": ("Informe Real de Inversión", "Se usa: hoja REAL <MES> <AÑO> o BASES <AÑO>"),
    "gral": ("GRAL Recepción", "No se necesita: es un resultado que esta app genera"),
    "provisorio": ("Planilla de la revisión provisoria", "No corresponde: se usa en la app de revisión de pagos provisorios"),
    "desconocido": ("No reconocido", "No tiene las pestañas esperadas"),
    "ilegible": ("No se pudo abrir", "Verifique que sea un .xlsx válido"),
}

st.title("Recepción real de pagos provisorios")
st.markdown(
    "**Etapa de gasto real.** Esta app no es la revisión de pagos provisorios "
    "(Estimado / Data Completa / Cronos): genera la hoja *Recepción PROVISORIOS* "
    "y el archivo *GRAL Recepción* a partir del archivo de cargas y del Informe Real del mes."
)

st.sidebar.header("Archivos del mes")
subidos = st.sidebar.file_uploader(
    "Suba los dos archivos (el orden no importa)", type=["xlsx"],
    accept_multiple_files=True, key="archivos",
    help="La app reconoce cada archivo por sus pestañas, no por su nombre.",
)

identificados = {}
filas_id = []
for f in subidos or []:
    b = f.getvalue()
    tipo = identificar_archivo(b, _hash(b))
    nombre, uso = DESCRIPCION_TIPO[tipo]
    filas_id.append({"Archivo subido": f.name, "Identificado como": nombre, "Uso": uso})
    if tipo in ("cargas", "informe"):
        if tipo in identificados:
            filas_id[-1]["Uso"] = "Duplicado: se usará el primero"
        else:
            identificados[tipo] = (f.name, b)

faltan = [t for t in ("cargas", "informe") if t not in identificados]
if faltan:
    st.subheader("Archivos necesarios")
    st.table(pd.DataFrame([
        {"Archivo": "Archivo de cargas de medios",
         "Ejemplo de nombre": "2025_CARGAS_MEDIOS_.xlsx",
         "Pestañas que se leen": "PROVISORIOS, Monitor Prov, Recepción PROVISORIOS, Aprobadores",
         "Estado": "Cargado" if "cargas" in identificados else "Falta"},
        {"Archivo": "Informe Real de Inversión",
         "Ejemplo de nombre": "INFORME REAL DE INVERSION CCU JULIO 2026.xlsx",
         "Pestañas que se leen": "REAL <MES> <AÑO> (o BASES <AÑO>)",
         "Estado": "Cargado" if "informe" in identificados else "Falta"},
    ]).set_index("Archivo"))
    st.caption(
        "El archivo GRAL RECEPCIÓN no se sube: es una de las dos salidas que la app "
        "genera, junto con la hoja Recepción PROVISORIOS."
    )
    if filas_id:
        st.subheader("Archivos subidos")
        st.table(pd.DataFrame(filas_id).set_index("Archivo subido"))
    st.stop()

with st.sidebar.expander("Archivos reconocidos", expanded=False):
    for t in ("cargas", "informe"):
        st.markdown(f"**{DESCRIPCION_TIPO[t][0]}:** {identificados[t][0]}")

b_cargas, b_informe = identificados["cargas"][1], identificados["informe"][1]

try:
    with st.spinner("Leyendo archivo de cargas..."):
        prov, monitor, hist, aprob = cargar_archivo_cargas(b_cargas, _hash(b_cargas))
    hojas = listar_hojas_informe(b_informe, _hash(b_informe))
except ErrorEstructura as e:
    st.error(str(e))
    st.stop()

if not hojas:
    st.error("El Informe Real no tiene hojas 'REAL <MES> <AÑO>' ni 'BASES <AÑO>'.")
    st.stop()

st.sidebar.header("Parámetros")
hoja = st.sidebar.selectbox("Hoja del Informe Real", hojas, index=0)
es_base = hoja.strip().upper().startswith(cfg.HOJA_BASES_REAL)

mes_sel = None
if es_base:
    mes_sel = st.sidebar.selectbox(
        "Mes a recepcionar", list(NOMBRE_MES), format_func=NOMBRE_MES.get,
        # El Informe Real se emite ~2 meses después (julio se emite en septiembre)
        index=(dt.date.today().month - 3) % 12,
    )

try:
    informe, mes, anio = cargar_informe(b_informe, _hash(b_informe), hoja, mes_sel)
except ErrorEstructura as e:
    st.error(str(e))
    st.stop()

if not es_base:
    st.sidebar.markdown(f"Mes detectado en la hoja: **{NOMBRE_MES.get(mes, '-')} {anio or ''}**")

excluir_ok = st.sidebar.checkbox(
    "Excluir líneas ya recepcionadas (con Documento material)", value=True,
    help="Si la Solped/posición ya aparece en Recepción PROVISORIOS con Documento material, el proceso está terminado.",
)
moneda_100 = st.sidebar.checkbox(
    "Líneas en USD se recepcionan al 100 %", value=True,
    help="El Informe Real viene en CLP; las líneas en USD (Meta/Google directo) se recepcionan por su valor provisorio.",
)
tol = st.sidebar.number_input(
    "Tolerancia de redondeo (pesos)", min_value=0, max_value=1000,
    value=int(cfg.TOLERANCIA_REDONDEO), step=1,
    help="Si |real - provisorio| es menor o igual a este valor, se mantiene el provisorio (100 %).",
)

# ---------------------------------------------------------------------------
# Proceso
# ---------------------------------------------------------------------------
clave_proceso = (_hash(b_cargas), _hash(b_informe), hoja, mes, excluir_ok, moneda_100, tol)
if st.session_state.get("clave_proceso") != clave_proceso:
    with st.spinner("Cruzando provisorios con Informe Real..."):
        st.session_state["resultado_base"] = procesar(
            prov, monitor, hist, aprob, informe, [mes],
            excluir_recepcionadas=excluir_ok, moneda_extranjera_100=moneda_100,
            tolerancia_redondeo=tol,
        )
    st.session_state["clave_proceso"] = clave_proceso
    st.session_state["ajustes"] = {}

res = aplicar_ajustes(st.session_state["resultado_base"], st.session_state.get("ajustes", {}))
r = res.resumen
mes_txt = f"{NOMBRE_MES.get(mes, '')} {anio or ''}".strip()

# ---------------------------------------------------------------------------
# Indicadores
# ---------------------------------------------------------------------------
st.subheader(f"Resultado: {mes_txt}")
c1, c2, c3, c4, c5 = st.columns(5)
c1.metric("Líneas generadas", r["lineas_generadas"], help=f"{r['excluidas']} excluidas de {r['lineas_mes']} del mes")
c2.metric("Cruzadas", r["cruzadas"])
c3.metric("A revisar", r["revisar"])
c4.metric("Sin cruce", r["sin_cruce"])
c5.metric("Moneda extranjera", r["moneda_extranjera"])

c1, c2, c3 = st.columns(3)
c1.metric("Provisorio CLP (J)", formato_pesos(r["monto_provisorio_clp"]))
c2.metric("Real CLP asignado (O)", formato_pesos(r["monto_real_clp"]),
          delta=formato_pesos(r["monto_real_clp"] - r["monto_provisorio_clp"]), delta_color="off")
c3.metric("Informe Real sin línea asociada", formato_pesos(r["monto_informe_sin_uso"]),
          help=f"{r['filas_informe_sin_uso']} filas del Informe Real no se asociaron a ninguna línea")

n_terminadas = int((res.excluidas.get("motivo_exclusion", pd.Series(dtype=str)) == "Ya recepcionada (tiene Documento material)").sum())
if n_terminadas and excluir_ok:
    st.info(
        f"{n_terminadas} líneas de {mes_txt} ya tienen Documento material en la Recepción "
        "PROVISORIOS histórica y no se vuelven a generar. Para regenerar el mes completo, "
        "desmarque 'Excluir líneas ya recepcionadas'."
    )

if r["lineas_generadas"] == 0:
    st.warning(
        "No hay líneas pendientes de recepción para este mes. Si el mes ya fue "
        "recepcionado, desmarque 'Excluir líneas ya recepcionadas' para regenerarlo."
    )

# ---------------------------------------------------------------------------
# Pestañas
# ---------------------------------------------------------------------------
t_rev, t_rec, t_gral, t_excl, t_inf, t_hom = st.tabs([
    "Revisión", "Recepción PROVISORIOS", "GRAL Recepción", "Excluidas",
    "Informe sin uso", "Homologación de marcas",
])

with t_rev:
    st.markdown(
        "Líneas sin cruce, con cruce a revisar o con observaciones. El monto real "
        "puede corregirse en la columna **Importe real** y se aplica a ambas salidas."
    )
    L = res.lineas
    mask = (
        L["estado_cruce"].isin(["Sin cruce", "Ajuste manual"])
        | L["estado_cruce"].str.contains("revisar")
        | (L["observacion"] != "")
    )
    rev = L.loc[mask, [
        "solped", "pos", "texto_breve", "texto_largo", "valor_prov", "monto_real",
        "estado_cruce", "observacion", "inf_pinta", "inf_vehiculo", "inf_tipo", "puntaje_cruce",
    ]].rename(columns={
        "solped": "Solped", "pos": "Pos.", "texto_breve": "Texto breve",
        "texto_largo": "Texto carga", "valor_prov": "Valor provisorio",
        "monto_real": "Importe real", "estado_cruce": "Estado",
        "observacion": "Observación", "inf_pinta": "Pinta", "inf_vehiculo": "Vehículo",
        "inf_tipo": "Tipo", "puntaje_cruce": "Puntaje",
    })
    if rev.empty:
        st.success("No hay líneas que requieran revisión.")
    else:
        editado = st.data_editor(
            rev, hide_index=True, width="stretch", key="editor_rev",
            disabled=[c for c in rev.columns if c != "Importe real"],
            column_config={
                "Valor provisorio": st.column_config.NumberColumn(format="%.0f"),
                "Importe real": st.column_config.NumberColumn(format="%.0f"),
            },
        )
        if st.button("Aplicar ajustes"):
            base = st.session_state["resultado_base"].lineas.set_index(["solped", "pos"])["monto_real"]
            ajustes = dict(st.session_state.get("ajustes", {}))
            for _, f in editado.iterrows():
                k = (f["Solped"], f["Pos."])
                nuevo = f["Importe real"]
                original = base.get(k)
                if pd.isna(nuevo) and pd.isna(original):
                    continue
                if pd.isna(nuevo) or pd.isna(original) or abs(float(nuevo) - float(original)) > 0.5:
                    ajustes[k] = None if pd.isna(nuevo) else float(nuevo)
            st.session_state["ajustes"] = ajustes
            st.rerun()
        if st.session_state.get("ajustes"):
            st.caption(f"Ajustes manuales aplicados: {len(st.session_state['ajustes'])}")

with t_rec:
    estados = sorted(res.recepcion["_Estado cruce"].unique())
    filtro = st.multiselect("Filtrar por estado de cruce", estados, default=estados)
    vista = res.recepcion[res.recepcion["_Estado cruce"].isin(filtro)]
    cols = cfg.COLUMNAS_RECEPCION[:-2] + ["SOX", "_Estado cruce", "_Pinta", "_Proveedor-Vehiculo", "_Burst", "_Monto Neto Real"]
    st.dataframe(vista[[c for c in cols if c]], hide_index=True, width="stretch")

with t_gral:
    resp = list(res.gral_por_responsable)
    sel = st.radio("Pestaña", ["Data"] + resp, horizontal=True)
    df = res.gral if sel == "Data" else res.gral_por_responsable[sel]
    st.dataframe(df[cfg.COLUMNAS_GRAL], hide_index=True, width="stretch")

with t_excl:
    if res.excluidas.empty:
        st.write("No hay líneas excluidas.")
    else:
        st.dataframe(
            res.excluidas[["solped", "pos", "pedido_txt", "texto_breve", "valor_prov",
                           "comentario", "motivo_exclusion"]],
            hide_index=True, width="stretch",
        )

with t_inf:
    st.markdown(
        "Filas del Informe Real del mes que no quedaron asociadas a ninguna línea "
        "provisoria (p. ej. Meta/Google directo, líneas aún no cargadas o ya recepcionadas)."
    )
    st.dataframe(
        res.informe_sin_uso[["fila_informe", "pinta", "medio", "vehiculo", "burst",
                             "tipo", "monto_real", "codigo_marca", "motivo_no_uso"]],
        hide_index=True, width="stretch",
    )

with t_hom:
    st.markdown(
        "Homologación de la Pinta del Informe Real con el código de marca del "
        "Elemento PEP. Las marcas sin homologar se agregan en "
        "`ALIAS_MARCA_A_CODIGO` (config.py)."
    )
    st.dataframe(res.homologacion, hide_index=True, width="stretch")

# ---------------------------------------------------------------------------
# Descargas
# ---------------------------------------------------------------------------
st.divider()
contexto = {"mes_txt": mes_txt, "hoja_informe": hoja, "fecha": dt.date.today().strftime("%d-%m-%Y")}
sufijo = f"{cfg.MESES.get(mes, ('', 'MES'))[1]}_{anio or ''}".strip("_")
d1, d2 = st.columns(2)
d1.download_button(
    "Descargar Recepción PROVISORIOS (.xlsx)",
    data=exportar_recepcion(res, contexto),
    file_name=f"Recepcion_PROVISORIOS_{sufijo}.xlsx",
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    width="stretch",
)
d2.download_button(
    "Descargar GRAL Recepción (.xlsx)",
    data=exportar_gral(res),
    file_name=f"GRAL_RECEPCION_{sufijo}.xlsx",
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    width="stretch",
)
