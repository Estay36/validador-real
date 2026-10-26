"""
Exportación a Excel con el mismo formato visual de las planillas originales.

  * Recepcion_PROVISORIOS_<MES>.xlsx
      - "Recepción PROVISORIOS": columnas A:AC idénticas a la pestaña del
        archivo de cargas (mismos colores de encabezado) y, desde AE, las
        columnas de trazabilidad del Informe Real.
      - Hojas de control: Resumen, Excluidas, Informe sin uso, Homologación.
  * GRAL_RECEPCION_<MES>.xlsx
      - "Data" + una pestaña por responsable (yu, ELO, RODRI, CONI).
"""
from __future__ import annotations

import io

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import config as cfg

_FUENTE = "Calibri"
_BORDE = Border(bottom=Side(style="thin", color="999999"))
_RELLENO_SIN_CRUCE = PatternFill("solid", fgColor="F4CCCC")
_RELLENO_REVISAR = PatternFill("solid", fgColor="FCE5CD")
_RELLENO_TRAZA = PatternFill("solid", fgColor="D9D9D9")
_RELLENO_AMARILLO = PatternFill("solid", fgColor="FFFF00")
_RELLENO_GRIS_GRAL = PatternFill("solid", fgColor="C0C0C0")

# Columnas del Informe Real con encabezado amarillo en el origen
_TRAZA_AMARILLAS = {"Pinta", "ID Medio", "Medio", "Proveedor-Vehiculo", "Burst", "Monto Neto Real"}

_FORMATO_NUM = {
    "Valor total": "#,##0.00", "Importe Contab. HES": "#,##0.00",
    "Desviacion por Posicion": "#,##0.00", "Porcentaje de Cumplimiento": "0.00",
    "Monto Neto Real": "#,##0", "Puntaje cruce": "0.0",
    "Por entregar (valor)": "#,##0.00", "Importe a Liberar": "#,##0.00",
    "Cantidad a Liberar": "#,##0.000", "Porcentaje a Liberar": "#,##0.00",
    "Proveedor fijo": "0", "Solicitud de pedido": "0", "Posición de pedido": "0",
    "Pedido": "0", "Cuenta de mayor": "0", "Servicio": "0", "Fila Informe Real": "0",
}


def _escribir_tabla(ws, df: pd.DataFrame, fila_ini=1, col_ini=1, rellenos=None,
                    ancho_max=45):
    rellenos = rellenos or {}
    for j, nombre in enumerate(df.columns):
        c = ws.cell(row=fila_ini, column=col_ini + j, value=nombre if nombre else None)
        c.font = Font(name=_FUENTE, bold=True, size=10)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.border = _BORDE
        if nombre in rellenos:
            c.fill = rellenos[nombre]
    for i, fila in enumerate(df.itertuples(index=False), start=1):
        for j, valor in enumerate(fila):
            if valor is not None and isinstance(valor, float) and pd.isna(valor):
                valor = None
            c = ws.cell(row=fila_ini + i, column=col_ini + j, value=valor)
            c.font = Font(name=_FUENTE, size=10)
            fmt = _FORMATO_NUM.get(df.columns[j])
            if fmt and isinstance(valor, (int, float)):
                c.number_format = fmt
            elif hasattr(valor, "year"):
                c.number_format = "dd-mm-yyyy"
    # Anchos de columna
    for j, nombre in enumerate(df.columns):
        letra = get_column_letter(col_ini + j)
        muestra = max((len(str(v)) for v in df.iloc[:200, j].tolist() if v is not None), default=0)
        ws.column_dimensions[letra].width = min(max(len(str(nombre)), muestra or 0, 6) + 2, ancho_max)
    ws.row_dimensions[fila_ini].height = 30


def _hoja_recepcion(wb, recepcion: pd.DataFrame):
    ws = wb.active
    ws.title = cfg.HOJA_RECEPCION_PROV
    plantilla = recepcion[cfg.COLUMNAS_RECEPCION]
    rellenos = {k: PatternFill("solid", fgColor=v) for k, v in cfg.COLOR_ENCABEZADO.items()}
    _escribir_tabla(ws, plantilla, rellenos=rellenos)

    traza = recepcion[[c for c in recepcion.columns if c.startswith("_")]].copy()
    traza.columns = [c[1:] for c in traza.columns]
    col_traza = len(cfg.COLUMNAS_RECEPCION) + 2          # deja AD en blanco
    rell_traza = {c: (_RELLENO_AMARILLO if c in _TRAZA_AMARILLAS else _RELLENO_TRAZA) for c in traza.columns}
    _escribir_tabla(ws, traza, col_ini=col_traza, rellenos=rell_traza)
    ws.column_dimensions[get_column_letter(col_traza - 1)].width = 3

    # Resaltado de filas que requieren atención (columna O)
    col_o = cfg.COLUMNAS_RECEPCION.index("Importe Contab. HES") + 1
    for i, estado in enumerate(recepcion["_Estado cruce"], start=2):
        if estado == "Sin cruce" or "sin cruce" in estado:
            ws.cell(row=i, column=col_o).fill = _RELLENO_SIN_CRUCE
        elif "revisar" in estado:
            ws.cell(row=i, column=col_o).fill = _RELLENO_REVISAR
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(col_traza + traza.shape[1] - 1)}{len(recepcion) + 1}"


def _hoja_simple(wb, titulo, df: pd.DataFrame):
    ws = wb.create_sheet(titulo)
    if df.empty:
        ws["A1"] = "Sin registros"
        return
    _escribir_tabla(ws, df)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(df.shape[1])}{len(df) + 1}"


def _hoja_resumen(wb, resumen: dict, contexto: dict):
    ws = wb.create_sheet("Resumen", 0)
    ws["A1"] = "Generación de Recepción PROVISORIOS con Informe Real"
    ws["A1"].font = Font(name=_FUENTE, bold=True, size=13)
    filas = [
        ("Mes procesado", contexto.get("mes_txt", "")),
        ("Hoja Informe Real", contexto.get("hoja_informe", "")),
        ("Fecha de generación", contexto.get("fecha", "")),
        ("", ""),
        ("Líneas provisorias del mes", resumen["lineas_mes"]),
        ("Excluidas (ver hoja Excluidas)", resumen["excluidas"]),
        ("Líneas generadas", resumen["lineas_generadas"]),
        ("  Cruzadas con Informe Real", resumen["cruzadas"]),
        ("    de ellas, a revisar", resumen["revisar"]),
        ("  Moneda extranjera (100 %)", resumen["moneda_extranjera"]),
        ("  Sin cruce (columna O vacía)", resumen["sin_cruce"]),
        ("", ""),
        ("Monto provisorio CLP (J)", resumen["monto_provisorio_clp"]),
        ("Monto real CLP asignado (O)", resumen["monto_real_clp"]),
        ("Monto total Informe Real", resumen["monto_informe_total"]),
        ("Monto Informe Real sin línea asociada", resumen["monto_informe_sin_uso"]),
        ("Filas Informe Real sin línea asociada", resumen["filas_informe_sin_uso"]),
    ]
    for i, (k, v) in enumerate(filas, start=3):
        ws.cell(row=i, column=1, value=k).font = Font(name=_FUENTE, size=10, bold=not k.startswith(" "))
        c = ws.cell(row=i, column=2, value=v)
        c.font = Font(name=_FUENTE, size=10)
        if isinstance(v, float):
            c.number_format = "#,##0"
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 24


def exportar_recepcion(res, contexto: dict) -> bytes:
    wb = Workbook()
    _hoja_recepcion(wb, res.recepcion)
    _hoja_resumen(wb, res.resumen, contexto)

    excl = res.excluidas[[
        "solped", "pos", "pedido_txt", "texto_breve", "valor_prov", "moneda",
        "comentario", "hoja_entrada_txt", "documento_material_hist", "motivo_exclusion",
    ]].rename(columns={
        "solped": "Solicitud de pedido", "pos": "Posición", "pedido_txt": "Pedido",
        "texto_breve": "Texto breve", "valor_prov": "Valor total", "moneda": "Moneda",
        "comentario": "Comentario PROVISORIOS", "hoja_entrada_txt": "Hoja de entrada (PROVISORIOS)",
        "documento_material_hist": "Documento material (histórico)",
        "motivo_exclusion": "Motivo exclusión",
    })
    _hoja_simple(wb, "Excluidas", excl)

    sin_uso = res.informe_sin_uso[[
        "fila_informe", "pinta", "id_medio", "medio", "vehiculo", "burst", "tipo",
        "monto_real", "codigo_marca", "motivo_no_uso",
    ]].rename(columns={
        "fila_informe": "Fila Informe Real", "pinta": "Pinta", "id_medio": "ID Medio",
        "medio": "Medio", "vehiculo": "Proveedor-Vehiculo", "burst": "Burst", "tipo": "Tipo",
        "monto_real": "Monto Neto Real", "codigo_marca": "Código marca PEP",
        "motivo_no_uso": "Motivo",
    })
    _hoja_simple(wb, "Informe sin uso", sin_uso)

    homol = res.homologacion.rename(columns={
        "pinta": "Pinta (Informe Real)", "pinta_norm": "Pinta normalizada",
        "codigo_marca": "Código marca PEP", "regla": "Regla aplicada",
    })
    _hoja_simple(wb, "Homologación marcas", homol)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def exportar_gral(res) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    rellenos = {c: _RELLENO_GRIS_GRAL for c in cfg.COLUMNAS_GRAL}
    _escribir_tabla(ws, res.gral[cfg.COLUMNAS_GRAL], rellenos=rellenos)
    for nombre, df in res.gral_por_responsable.items():
        w = wb.create_sheet(str(nombre)[:31])
        _escribir_tabla(w, df[cfg.COLUMNAS_GRAL], rellenos=rellenos)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
